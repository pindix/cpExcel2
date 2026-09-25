# -*- coding: utf-8 -*-
"""
Converte medicamentos.xlsx -> medicamentos.js (JSON embutido numa variável global).

Porque .js em vez de .json?
  - Funciona com duplo-clique no index.html (file://) sem servidor local.
  - O fetch() de ficheiros .json é bloqueado por CORS quando abres via file://.
  - Um <script src="medicamentos.js"> carrega sempre, mesmo offline.
"""

import json
import os
import sys
from openpyxl import load_workbook

FICHEIRO_EXCEL = "medicamentos.xlsx"
FICHEIRO_SAIDA = "medicamentos.js"

def formatar_folha(ws):
    """Reproduz exatamente a lógica do formatarFolha() do JS."""
    # Ler tudo para lista de listas
    linhas = []
    for row in ws.iter_rows(values_only=True):
        linhas.append(list(row))

    # Encontrar a linha de cabeçalho (primeira célula == "nome")
    idx_cabecalho = -1
    for i, linha in enumerate(linhas):
        if linha and linha[0] is not None and str(linha[0]).strip().lower() == "nome":
            idx_cabecalho = i
            break
    if idx_cabecalho == -1:
        return []

    cabecalho = [str(c).strip().lower() if c is not None else "" for c in linhas[idx_cabecalho]]
    dados = linhas[idx_cabecalho + 1:]

    registos = []
    for linha in dados:
        obj = {}
        for i, col in enumerate(cabecalho):
            if col == "":
                continue
            valor = linha[i] if i < len(linha) else ""
            if valor is None:
                valor = ""
            # Mesma regra do JS: coluna "nome" com "|" vira lista
            if col == "nome" and isinstance(valor, str) and "|" in valor:
                obj[col] = [s.strip() for s in valor.split("|")]
            else:
                obj[col] = valor
        # Ignorar linhas totalmente vazias
        if any(v not in ("", None) for v in obj.values()):
            registos.append(obj)
    return registos


def main():
    if not os.path.exists(FICHEIRO_EXCEL):
        print(f"❌ Não encontrei '{FICHEIRO_EXCEL}' nesta pasta: {os.getcwd()}")
        sys.exit(1)

    print(f"📖 A ler '{FICHEIRO_EXCEL}'...")
    wb = load_workbook(FICHEIRO_EXCEL, data_only=True)

    banco = {}
    for nome_folha in wb.sheetnames:
        ws = wb[nome_folha]
        chave = nome_folha.strip().lower()
        registos = formatar_folha(ws)
        banco[chave] = registos
        print(f"   • Folha '{nome_folha}' → {len(registos)} registos")

    # Escrever como variável JS global (mais seguro que .json via file://)
    with open(FICHEIRO_SAIDA, "w", encoding="utf-8") as f:
        f.write("// Gerado automaticamente por excel_para_json.py\n")
        f.write("// NÃO editar à mão.\n")
        f.write("window.BANCO_DADOS = ")
        json.dump(banco, f, ensure_ascii=False, indent=2)
        f.write(";\n")

    tamanho_kb = os.path.getsize(FICHEIRO_SAIDA) / 1024
    print(f"\n✅ Criado '{FICHEIRO_SAIDA}' ({tamanho_kb:.1f} KB)")
    print("   Agora basta abrir o index.html — já funciona offline.")


if __name__ == "__main__":
    main()